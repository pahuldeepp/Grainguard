"""Normalize the Mendeley stored-wheat insect workbook to the proxy CSV contract."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import sys

SOURCE_URL = "https://data.mendeley.com/datasets/53ckhtd8gj/1"
EXPECTED = {"Silo", "Time_Index", "Temperature", "Humidity", "Count"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(source: Path, output: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("install ml/real_data/proxy-requirements.txt to read the source workbook") from exc
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Data"]:
            raise ValueError("expected exactly one source worksheet named Data")
        sheet = workbook["Data"]
        headers = [str(value).strip() if value is not None else "" for value in
                   next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if not EXPECTED.issubset(headers):
            raise ValueError(f"source workbook is missing required columns: {sorted(EXPECTED - set(headers))}")
        columns = {name: headers.index(name) for name in EXPECTED}
        output.parent.mkdir(parents=True, exist_ok=True)
        count = 0
        with output.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(["silo_id", "time_index", "temperature_c", "grain_moisture_percent", "insect_count"])
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                if row[columns["Silo"]] is None:
                    continue
                try:
                    # Source column is named Humidity; source file and paper
                    # describe this as grain moisture, not ambient RH.
                    silo = int(row[columns["Silo"]])
                    time_index = int(row[columns["Time_Index"]])
                    temperature = float(row[columns["Temperature"]])
                    moisture = float(row[columns["Humidity"]])
                    insect_count = int(row[columns["Count"]])
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"source row {row_number}: invalid value") from exc
                if silo < 1 or time_index < 1 or insect_count < 0:
                    raise ValueError(f"source row {row_number}: value out of range")
                writer.writerow([f"silo-{silo}", time_index, temperature, moisture, insect_count])
                count += 1
    finally:
        workbook.close()
    record = {
        "schema_version": "greenguard-public-proxy-source-v1",
        "dataset_id": "mendeley-stored-wheat-insect-activity-v1",
        "source_url": SOURCE_URL,
        "source_title": "Integrated Modeling of Stored-Grain Insect Dynamics and Infestation Risk in Wheat Silos",
        "license": "CC BY 4.0",
        "source_file": source.name,
        "source_file_sha256": sha256(source),
        "normalized_csv_sha256": sha256(output),
        "normalizer_sha256": sha256(Path(__file__)),
        "rows": count,
        "field_mapping": {"Humidity": "grain_moisture_percent (not relative humidity)",
                          "Temperature": "temperature_c", "Count": "observed insect count"},
        "target_boundary": "insect activity only; not spoilage/mildew/mycotoxin",
        "production_eligible": False,
        "promotion_eligible": False,
    }
    sidecar = output.with_suffix(output.suffix + ".source.json")
    sidecar.write_text(json.dumps(record, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return record


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, type=Path, help="Mendeley integrated workbook; keep outside Git")
    parser.add_argument("--output", required=True, type=Path, help="normalized CSV; keep outside Git")
    args = parser.parse_args(argv)
    try:
        print(json.dumps(normalize(args.xlsx, args.output), sort_keys=True, indent=2))
    except (OSError, RuntimeError, ValueError, StopIteration) as exc:
        print(json.dumps({"status": "rejected", "error": str(exc)}, sort_keys=True), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
